#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""validate_deck.py - deterministic PPTX brand-compliance validator.

Walks every shape on every slide of a .pptx (recursing into groups, tables and
charts) and reports brand violations against a brand profile plus the shared
deck grammar.

    validate_deck.py <file.pptx> [--brand ID] [--format json|human]
                     [--ir deck.json] [--max-warn N]

Output is FROZEN CONTRACT C:

    {"target":str,"brand":str,"kind":"deck","pass":bool,
     "counts":{"error":int,"warn":int,"info":int},
     "violations":[{"id","severity","where","found","expected","rule","fix"}]}

Exit codes
    0   no errors
    2   one or more errors (or the --max-warn budget was blown)
    1   internal failure (bad arguments, unreadable file, unknown brand)

Zero third-party dependencies beyond python-pptx and pillow.
Python 3.9 compatible.
"""

from __future__ import absolute_import

import argparse
import hashlib
import json
import os
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from typing import Any, Dict, List, Optional, Sequence, Tuple

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

from lib import brandlib as bl  # noqa: E402

try:
    from pptx import Presentation
    from pptx.oxml.ns import qn
except ImportError as _exc:  # pragma: no cover - environment problem
    sys.stderr.write("validate_deck.py requires python-pptx: %s\n" % _exc)
    raise SystemExit(1)


# ---------------------------------------------------------------------------
# constants
# ---------------------------------------------------------------------------

KIND = "deck"
DEFAULT_BRAND = "channelplay"
CUSTOM_PROP_NAME = "brand-studio.brand"

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"

_CUSTOM_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/custom-properties}"

#: Slack, in inches, when matching a shape to a grammar region by geometry.
REGION_TOL_IN = 0.22
#: Slack, in inches, for the off-canvas test (template art routinely sits at -0.01).
CANVAS_TOL_IN = 0.03
#: Slack, in points, when matching a run size to the brand type scale.
SCALE_TOL_PT = 0.6
#: A colour this close to #000000 counts as black text.
BLACK_DELTA_E = 6.0
#: Fraction of a text box that must sit inside a shape for it to count as its background.
CONTAINMENT_RATIO = 0.98
#: Rendered/declared logo aspect may differ by this fraction before it is "distorted".
LOGO_ASPECT_TOL = 0.02

_IMAGE_EXT_RE = re.compile(r"\.(png|jpe?g|gif|bmp|tiff?|emf|wmf|svg|webp)$", re.I)
_FILENAMEISH_RE = re.compile(r"^[\w\-. ]+\.(png|jpe?g|gif|bmp|tiff?|emf|wmf|svg|webp)$", re.I)

#: The handful of DrawingML preset colours worth resolving.
PRESET_COLORS = {
    "black": "#000000", "white": "#FFFFFF", "red": "#FF0000", "green": "#008000",
    "blue": "#0000FF", "yellow": "#FFFF00", "gray": "#808080", "grey": "#808080",
    "darkGray": "#A9A9A9", "lightGray": "#D3D3D3", "cyan": "#00FFFF",
    "magenta": "#FF00FF", "orange": "#FFA500", "purple": "#800080",
}

#: MSO scheme-colour token -> a:clrScheme child tag.
SCHEME_ALIASES = {
    "tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2",
}


# ---------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------

def _local(tag):
    # type: (Any) -> str
    """Local name of an lxml/ElementTree tag."""
    t = str(tag)
    return t.rsplit("}", 1)[-1]


def _num(el, attr, default=0):
    # type: (Any, str, int) -> int
    if el is None:
        return default
    v = el.get(attr)
    if v is None:
        return default
    try:
        return int(v)
    except (TypeError, ValueError):
        try:
            return int(float(v))
        except (TypeError, ValueError):
            return default


def _clamp(v, lo, hi):
    # type: (float, float, float) -> float
    return lo if v < lo else (hi if v > hi else v)


def _rgb_to_hsl(h):
    # type: (str) -> Tuple[float, float, float]
    r, g, b = [c / 255.0 for c in bl.hex_to_rgb(h)]
    mx, mn = max(r, g, b), min(r, g, b)
    lum = (mx + mn) / 2.0
    if abs(mx - mn) < 1e-12:
        return (0.0, 0.0, lum)
    d = mx - mn
    sat = d / (2.0 - mx - mn) if lum > 0.5 else d / (mx + mn)
    if mx == r:
        hue = ((g - b) / d) % 6.0
    elif mx == g:
        hue = (b - r) / d + 2.0
    else:
        hue = (r - g) / d + 4.0
    return (hue * 60.0, sat, lum)


def _hsl_to_rgb_hex(hue, sat, lum):
    # type: (float, float, float) -> str
    hue = hue % 360.0
    sat = _clamp(sat, 0.0, 1.0)
    lum = _clamp(lum, 0.0, 1.0)
    c = (1.0 - abs(2.0 * lum - 1.0)) * sat
    x = c * (1.0 - abs((hue / 60.0) % 2.0 - 1.0))
    m = lum - c / 2.0
    if hue < 60:
        rgb = (c, x, 0.0)
    elif hue < 120:
        rgb = (x, c, 0.0)
    elif hue < 180:
        rgb = (0.0, c, x)
    elif hue < 240:
        rgb = (0.0, x, c)
    elif hue < 300:
        rgb = (x, 0.0, c)
    else:
        rgb = (c, 0.0, x)
    return bl.rgb_to_hex([(v + m) * 255.0 for v in rgb])


def _apply_shade(hex_value, factor):
    # type: (str, float) -> str
    r, g, b = bl.hex_to_rgb(hex_value)
    return bl.rgb_to_hex([r * factor, g * factor, b * factor])


def _apply_tint(hex_value, factor):
    # type: (str, float) -> str
    r, g, b = bl.hex_to_rgb(hex_value)
    return bl.rgb_to_hex([c * factor + 255.0 * (1.0 - factor) for c in (r, g, b)])


# ---------------------------------------------------------------------------
# COLOUR EXTRACTION
# ---------------------------------------------------------------------------

class ThemeColors(object):
    """Resolves ``a:schemeClr`` tokens for one slide master."""

    def __init__(self, mapping=None, clr_map=None):
        # type: (Optional[Dict[str, str]], Optional[Dict[str, str]]) -> None
        self.mapping = mapping or {}
        self.clr_map = clr_map or {}

    def get(self, token):
        # type: (str) -> Optional[str]
        key = str(token or "")
        if key == "phClr":
            return None
        key = self.clr_map.get(key, key)
        key = SCHEME_ALIASES.get(key, key)
        return self.mapping.get(key)


def _theme_for_slide(slide, cache):
    # type: (Any, Dict[Any, ThemeColors]) -> ThemeColors
    """Build (and cache) the colour scheme backing a slide."""
    try:
        master = slide.slide_layout.slide_master
    except Exception:
        return ThemeColors()
    key = id(master)
    if key in cache:
        return cache[key]

    mapping = {}  # type: Dict[str, str]
    try:
        theme_part = None
        for rel in master.part.rels.values():
            if rel.is_external:
                continue
            if str(rel.reltype).endswith("/theme"):
                theme_part = rel.target_part
                break
        if theme_part is not None:
            root = ET.fromstring(theme_part.blob)
            scheme = root.find("./%sthemeElements/%sclrScheme" % (_A, _A))
            if scheme is not None:
                for child in list(scheme):
                    name = _local(child.tag)
                    hexed = _color_from_element_tree(child)
                    if hexed:
                        mapping[name] = hexed
    except Exception:
        mapping = {}

    clr_map = {}  # type: Dict[str, str]
    try:
        el = master._element.find(qn("p:clrMap"))
        if el is not None:
            for attr, value in el.attrib.items():
                clr_map[_local(attr)] = str(value)
    except Exception:
        clr_map = {}

    theme = ThemeColors(mapping, clr_map)
    cache[key] = theme
    return theme


def _color_from_element_tree(parent):
    # type: (Any) -> Optional[str]
    """First resolvable colour child of an ElementTree node (theme parsing only)."""
    for child in list(parent):
        name = _local(child.tag)
        if name == "srgbClr":
            val = child.get("val")
            if val and bl.is_hex(val):
                return bl.normalize_hex(val)
        elif name == "sysClr":
            val = child.get("lastClr")
            if val and bl.is_hex(val):
                return bl.normalize_hex(val)
    return None


def _color_hex(clr_el, theme):
    # type: (Any, ThemeColors) -> Optional[str]
    """Resolve one DrawingML colour element to '#RRGGBB', honouring modifiers."""
    if clr_el is None:
        return None
    name = _local(clr_el.tag)
    base = None  # type: Optional[str]
    if name == "srgbClr":
        val = clr_el.get("val")
        if val and bl.is_hex(val):
            base = bl.normalize_hex(val)
    elif name == "schemeClr":
        base = theme.get(clr_el.get("val") or "")
    elif name == "sysClr":
        val = clr_el.get("lastClr")
        if val and bl.is_hex(val):
            base = bl.normalize_hex(val)
    elif name == "prstClr":
        base = PRESET_COLORS.get(str(clr_el.get("val") or ""))
    elif name == "scrgbClr":
        try:
            base = bl.rgb_to_hex([
                float(clr_el.get("r", 0)) / 100000.0 * 255.0,
                float(clr_el.get("g", 0)) / 100000.0 * 255.0,
                float(clr_el.get("b", 0)) / 100000.0 * 255.0,
            ])
        except (TypeError, ValueError):
            base = None
    elif name == "hslClr":
        try:
            base = _hsl_to_rgb_hex(float(clr_el.get("hue", 0)) / 60000.0,
                                   float(clr_el.get("sat", 0)) / 100000.0,
                                   float(clr_el.get("lum", 0)) / 100000.0)
        except (TypeError, ValueError):
            base = None
    if base is None:
        return None

    for mod in list(clr_el):
        mname = _local(mod.tag)
        try:
            mval = float(mod.get("val")) if mod.get("val") is not None else None
        except (TypeError, ValueError):
            mval = None
        if mval is None:
            continue
        if mname == "lumMod":
            hue, sat, lum = _rgb_to_hsl(base)
            base = _hsl_to_rgb_hex(hue, sat, lum * (mval / 100000.0))
        elif mname == "lumOff":
            hue, sat, lum = _rgb_to_hsl(base)
            base = _hsl_to_rgb_hex(hue, sat, lum + (mval / 100000.0))
        elif mname == "shade":
            base = _apply_shade(base, mval / 100000.0)
        elif mname == "tint":
            base = _apply_tint(base, mval / 100000.0)
        elif mname == "satMod":
            hue, sat, lum = _rgb_to_hsl(base)
            base = _hsl_to_rgb_hex(hue, sat * (mval / 100000.0), lum)
    return base


def _first_color_child(parent, theme):
    # type: (Any, ThemeColors) -> Optional[str]
    if parent is None:
        return None
    for child in list(parent):
        hexed = _color_hex(child, theme)
        if hexed:
            return hexed
    return None


def fill_info(container, theme):
    # type: (Any, ThemeColors) -> Dict[str, Any]
    """Describe the fill declared on an spPr / rPr / tcPr element.

    Returns {"kind": None|"solid"|"gradient"|"none"|"picture"|"pattern"|"group",
             "colors": ['#RRGGBB', ...]}.  ``kind`` is None when nothing is
    declared, i.e. the fill is inherited from the layout, master or theme.
    """
    out = {"kind": None, "colors": []}  # type: Dict[str, Any]
    if container is None:
        return out

    solid = container.find(qn("a:solidFill"))
    if solid is not None:
        hexed = _first_color_child(solid, theme)
        out["kind"] = "solid"
        out["colors"] = [hexed] if hexed else []
        return out

    grad = container.find(qn("a:gradFill"))
    if grad is not None:
        colors = []
        lst = grad.find(qn("a:gsLst"))
        if lst is not None:
            for gs in lst.findall(qn("a:gs")):
                hexed = _first_color_child(gs, theme)
                if hexed:
                    colors.append(hexed)
        out["kind"] = "gradient"
        out["colors"] = colors
        return out

    if container.find(qn("a:noFill")) is not None:
        out["kind"] = "none"
        return out
    if container.find(qn("a:blipFill")) is not None:
        out["kind"] = "picture"
        return out
    patt = container.find(qn("a:pattFill"))
    if patt is not None:
        colors = []
        for tag in ("a:fgClr", "a:bgClr"):
            hexed = _first_color_child(patt.find(qn(tag)), theme)
            if hexed:
                colors.append(hexed)
        out["kind"] = "pattern"
        out["colors"] = colors
        return out
    if container.find(qn("a:grpFill")) is not None:
        out["kind"] = "group"
    return out


# ---------------------------------------------------------------------------
# SHAPE FLATTENING
# ---------------------------------------------------------------------------

class ShapeRec(object):
    """One leaf shape with absolute geometry, in inches."""

    __slots__ = ("shape", "el", "name", "tag", "left", "top", "width", "height",
                 "z", "path", "text", "runs", "fill", "line", "is_picture",
                 "is_table", "is_chart", "role", "semantic", "logo_variant",
                 "logo_conf", "in_group")

    def __init__(self):
        self.shape = None
        self.el = None
        self.name = ""
        self.tag = ""
        self.left = 0.0
        self.top = 0.0
        self.width = 0.0
        self.height = 0.0
        self.z = 0
        self.path = ""
        self.text = ""
        self.runs = []          # type: List[Dict[str, Any]]
        self.fill = {"kind": None, "colors": []}
        self.line = {"kind": None, "colors": []}
        self.is_picture = False
        self.is_table = False
        self.is_chart = False
        self.role = None        # type: Optional[str]
        self.semantic = None    # type: Optional[str]
        self.logo_variant = None  # type: Optional[str]
        self.logo_conf = None     # type: Optional[str]
        self.in_group = False

    @property
    def rect(self):
        # type: () -> Tuple[float, float, float, float]
        return (self.left, self.top, self.width, self.height)

    @property
    def area(self):
        # type: () -> float
        return max(0.0, self.width) * max(0.0, self.height)


def _identity_xform(x, y, w, h):
    # type: (float, float, float, float) -> Tuple[float, float, float, float]
    return (x, y, w, h)


def _group_xform(group_el, parent_xform):
    """Return a child->absolute EMU transform for a p:grpSp element."""
    xfrm = None
    grp_pr = group_el.find(qn("p:grpSpPr"))
    if grp_pr is not None:
        xfrm = grp_pr.find(qn("a:xfrm"))
    if xfrm is None:
        return parent_xform

    off = xfrm.find(qn("a:off"))
    ext = xfrm.find(qn("a:ext"))
    ch_off = xfrm.find(qn("a:chOff"))
    ch_ext = xfrm.find(qn("a:chExt"))

    gx, gy = _num(off, "x"), _num(off, "y")
    gw, gh = _num(ext, "cx"), _num(ext, "cy")
    ax, ay, aw, ah = parent_xform(gx, gy, gw, gh)

    cox, coy = _num(ch_off, "x"), _num(ch_off, "y")
    cex, cey = _num(ch_ext, "cx", 0), _num(ch_ext, "cy", 0)
    sx = (float(aw) / float(cex)) if cex else 1.0
    sy = (float(ah) / float(cey)) if cey else 1.0

    def xform(x, y, w, h):
        return (ax + (x - cox) * sx, ay + (y - coy) * sy, w * sx, h * sy)

    return xform


def flatten_shapes(container, xform, out, counter, prefix, in_group=False):
    # type: (Any, Any, List[ShapeRec], List[int], str, bool) -> None
    """Append every leaf shape, with absolute geometry, in paint (z) order."""
    for shape in container:
        el = shape._element
        tag = _local(el.tag)

        if tag == "grpSp":
            child_xform = _group_xform(el, xform)
            gname = _shape_name(el) or "group"
            try:
                children = shape.shapes
            except Exception:
                children = []
            flatten_shapes(children, child_xform, out, counter,
                           "%s%s/" % (prefix, gname), True)
            continue

        rec = ShapeRec()
        rec.shape = shape
        rec.el = el
        rec.tag = tag
        rec.name = _shape_name(el) or tag
        rec.in_group = in_group
        left = shape.left if shape.left is not None else 0
        top = shape.top if shape.top is not None else 0
        width = shape.width if shape.width is not None else 0
        height = shape.height if shape.height is not None else 0
        ax, ay, aw, ah = xform(int(left), int(top), int(width), int(height))
        rec.left = bl.emu_to_in(ax)
        rec.top = bl.emu_to_in(ay)
        rec.width = bl.emu_to_in(aw)
        rec.height = bl.emu_to_in(ah)
        rec.z = counter[0]
        counter[0] += 1
        rec.path = "%s%s" % (prefix, rec.name)
        rec.is_picture = (tag == "pic")
        rec.is_table = bool(getattr(shape, "has_table", False))
        rec.is_chart = bool(getattr(shape, "has_chart", False))
        out.append(rec)


def _shape_name(el):
    # type: (Any) -> str
    try:
        found = el.xpath("./*/p:cNvPr")
        if found:
            return str(found[0].get("name") or "")
    except Exception:
        pass
    return ""


def _shape_descr(el):
    # type: (Any) -> str
    try:
        found = el.xpath("./*/p:cNvPr")
        if found:
            return str(found[0].get("descr") or "")
    except Exception:
        pass
    return ""


# ---------------------------------------------------------------------------
# TEXT EXTRACTION
# ---------------------------------------------------------------------------

def _def_rpr(p_el):
    # type: (Any) -> Any
    p_pr = p_el.find(qn("a:pPr"))
    if p_pr is None:
        return None
    return p_pr.find(qn("a:defRPr"))


def _lst_style_rpr(body_el):
    # type: (Any) -> Any
    if body_el is None:
        return None
    lst = body_el.find(qn("a:lstStyle"))
    if lst is None:
        return None
    lvl = lst.find(qn("a:lvl1pPr"))
    if lvl is None:
        return None
    return lvl.find(qn("a:defRPr"))


def _typeface(rpr):
    # type: (Any) -> Optional[str]
    if rpr is None:
        return None
    latin = rpr.find(qn("a:latin"))
    if latin is None:
        return None
    tf = latin.get("typeface")
    return str(tf) if tf else None


def _size_pt(rpr):
    # type: (Any) -> Optional[float]
    if rpr is None:
        return None
    sz = rpr.get("sz")
    if sz is None:
        return None
    try:
        return float(sz) / 100.0
    except (TypeError, ValueError):
        return None


def _bool_attr(rpr, attr):
    # type: (Any, str) -> Optional[bool]
    if rpr is None:
        return None
    v = rpr.get(attr)
    if v is None:
        return None
    return str(v) in ("1", "true", "True")


def _spc(rpr):
    # type: (Any) -> Optional[int]
    if rpr is None:
        return None
    v = rpr.get("spc")
    if v is None:
        return None
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def _is_bulleted(p_el):
    # type: (Any) -> bool
    p_pr = p_el.find(qn("a:pPr"))
    if p_pr is None:
        return False
    if p_pr.find(qn("a:buNone")) is not None:
        return False
    for tag in ("a:buChar", "a:buAutoNum", "a:buBlip"):
        if p_pr.find(qn(tag)) is not None:
            return True
    return False


def extract_paragraphs(text_frame, theme):
    # type: (Any, ThemeColors) -> List[Dict[str, Any]]
    """Flatten a text frame into paragraph records carrying resolved run data."""
    paras = []  # type: List[Dict[str, Any]]
    try:
        body_el = text_frame._txBody
    except Exception:
        body_el = None
    lst_rpr = _lst_style_rpr(body_el)

    try:
        paragraph_list = list(text_frame.paragraphs)
    except Exception:
        paragraph_list = []

    for pi, para in enumerate(paragraph_list):
        p_el = para._p
        d_rpr = _def_rpr(p_el)
        runs = []  # type: List[Dict[str, Any]]
        for ri, run in enumerate(para.runs):
            r_el = run._r
            rpr = r_el.find(qn("a:rPr"))

            family = _typeface(rpr)
            inherited_family = family is None
            if family is None:
                family = _typeface(d_rpr)
            if family is None:
                family = _typeface(lst_rpr)

            size = _size_pt(rpr)
            if size is None:
                size = _size_pt(d_rpr)
            if size is None:
                size = _size_pt(lst_rpr)

            bold = _bool_attr(rpr, "b")
            if bold is None:
                bold = _bool_attr(d_rpr, "b")
            italic = _bool_attr(rpr, "i")
            if italic is None:
                italic = _bool_attr(d_rpr, "i")

            spc = _spc(rpr)
            if spc is None:
                spc = _spc(d_rpr)

            fill = fill_info(rpr, theme)
            if fill["kind"] is None:
                fill = fill_info(d_rpr, theme)
            if fill["kind"] is None:
                fill = fill_info(lst_rpr, theme)

            runs.append({
                "index": ri,
                "text": run.text or "",
                "family": family,
                "family_inherited": bool(inherited_family),
                "size": size,
                "bold": bool(bold),
                "italic": bool(italic),
                "spc": spc,
                "fill": fill,
            })

        text = (para.text or "").replace("\v", "\n")
        paras.append({
            "index": pi,
            "text": text,
            "runs": runs,
            "bulleted": _is_bulleted(p_el),
        })
    return paras


def paragraph_size(para, fallback):
    # type: (Dict[str, Any], float) -> float
    sizes = [r["size"] for r in para["runs"] if r["size"]]
    if sizes:
        return float(max(sizes))
    return float(fallback)


def paragraph_family(para, fallback):
    # type: (Dict[str, Any], Optional[str]) -> Optional[str]
    for r in para["runs"]:
        if r["family"]:
            return r["family"]
    return fallback


# ---------------------------------------------------------------------------
# GRAMMAR / ARCHETYPE
# ---------------------------------------------------------------------------

class Grammar(object):
    """Convenience view over deck-grammar.json."""

    def __init__(self, data):
        # type: (Dict[str, Any]) -> None
        self.data = data
        canvas = data.get("canvas") or {}
        self.canvas_w = float(canvas.get("w", 13.333))
        self.canvas_h = float(canvas.get("h", 7.5))
        self.grid = data.get("grid") or {}
        self.chrome = data.get("chrome") or {}
        self.rules = data.get("deckRules") or {}
        self.archetypes = {}  # type: Dict[str, Dict[str, Any]]
        self.order = []       # type: List[str]
        for arch in data.get("archetypes") or []:
            aid = str(arch.get("id") or "")
            if aid:
                self.archetypes[aid] = arch
                self.order.append(aid)

    def regions(self, archetype_id):
        # type: (Optional[str]) -> List[Dict[str, Any]]
        """Absolute region rects for an archetype, chrome included when it applies."""
        out = []  # type: List[Dict[str, Any]]
        arch = self.archetypes.get(archetype_id or "")
        if arch is None:
            return out
        if arch.get("chrome"):
            for name, spec in self.chrome.items():
                if isinstance(spec, dict) and "x" in spec:
                    out.append(self._region(name, spec))
        for name, spec in (arch.get("regions") or {}).items():
            if isinstance(spec, dict) and "x" in spec:
                out.append(self._region(name, spec))
        return out

    @staticmethod
    def _region(name, spec):
        # type: (str, Dict[str, Any]) -> Dict[str, Any]
        return {
            "name": name,
            "x": float(spec.get("x", 0.0)),
            "y": float(spec.get("y", 0.0)),
            "w": float(spec.get("w", 0.0)),
            "h": float(spec.get("h", 0.0)),
            "kind": str(spec.get("kind") or ""),
            "role": str(spec.get("role") or ""),
            "optional": bool(spec.get("optional")),
        }

    def column_xs(self):
        # type: () -> List[float]
        left = float(self.grid.get("contentLeft", 0.869))
        cw = float(self.grid.get("columnWidth", 0.7325))
        gap = float(self.grid.get("columnGap", 0.28))
        cols = int(self.grid.get("columns", 12))
        return [left + n * (cw + gap) for n in range(max(0, cols))]

    def set_xs(self, archetype_id):
        # type: (Optional[str]) -> List[float]
        """Left edges of the items inside an archetype's grid/stat/step/column sets."""
        out = []  # type: List[float]
        arch = self.archetypes.get(archetype_id or "")
        if arch is None:
            return out
        for spec in (arch.get("regions") or {}).values():
            if not isinstance(spec, dict) or "x" not in spec:
                continue
            kind = str(spec.get("kind") or "")
            if kind not in ("grid", "statSet", "stepSet", "stageSet", "columnSet",
                            "photoSet"):
                continue
            try:
                x = float(spec.get("x", 0.0))
                w = float(spec.get("w", 0.0))
            except (TypeError, ValueError):
                continue
            cols = spec.get("cols")
            if cols is None:
                cols = spec.get("count", 1)
            try:
                cols = int(cols)
            except (TypeError, ValueError):
                cols = 1
            gap = spec.get("colGap")
            if gap is None:
                gap = spec.get("gap", 0.0)
            try:
                gap = float(gap)
            except (TypeError, ValueError):
                gap = 0.0
            if cols < 1 or w <= 0:
                continue
            item_w = (w - gap * (cols - 1)) / float(cols)
            for n in range(cols):
                out.append(x + n * (item_w + gap))
        return out

    def spec_dxs(self, archetype_id):
        # type: (Optional[str]) -> List[float]
        """Every horizontal offset declared by an archetype's item/cell/row specs."""
        out = []  # type: List[float]
        arch = self.archetypes.get(archetype_id or "")
        if arch is None:
            return out
        for key, spec in arch.items():
            if not str(key).endswith("Spec") or not isinstance(spec, dict):
                continue
            for part in spec.values():
                if isinstance(part, dict) and "dx" in part:
                    try:
                        out.append(float(part["dx"]))
                    except (TypeError, ValueError):
                        continue
        return sorted(set(out))

    def rule(self, key, default=None):
        # type: (str, Any) -> Any
        return self.rules.get(key, default)


def infer_archetype(recs, grammar, slide_index, slide_count):
    # type: (List[ShapeRec], Grammar, int, int) -> Optional[str]
    """Best-guess archetype from shape geometry. Deterministic, never raises."""
    opener = str(grammar.rule("mustOpenWith", "cover"))
    best_id = None
    best_score = 0.0
    for aid in grammar.order:
        if aid == opener and slide_index != 1:
            continue  # a cover can only ever be slide 1
        regions = [r for r in grammar.regions(aid) if r["kind"] in ("text", "image", "visual",
                                                                    "rect", "grid", "table",
                                                                    "statSet", "stepSet",
                                                                    "stageSet", "columnSet",
                                                                    "photoSet", "rowList",
                                                                    "numberedList")]
        required = [r for r in regions if not r["optional"]]
        pool = required or regions
        if not pool:
            continue
        matched = 0
        for region in pool:
            for rec in recs:
                if (abs(rec.left - region["x"]) <= REGION_TOL_IN
                        and abs(rec.top - region["y"]) <= REGION_TOL_IN):
                    matched += 1
                    break
        score = float(matched) / float(len(pool))
        if score > best_score:
            best_score = score
            best_id = aid
    del slide_count
    if best_score >= 0.5:
        return best_id
    return None


# ---------------------------------------------------------------------------
# CONTEXT
# ---------------------------------------------------------------------------

class SlideCtx(object):
    __slots__ = ("index", "slide", "recs", "archetype", "archetype_source",
                 "ir", "bg_hex", "theme", "regions", "_ir_roles")

    def __init__(self):
        self.index = 0
        self.slide = None
        self.recs = []            # type: List[ShapeRec]
        self.archetype = None     # type: Optional[str]
        self.archetype_source = "none"
        self.ir = None            # type: Optional[Dict[str, Any]]
        self.bg_hex = "#FFFFFF"
        self.theme = ThemeColors()
        self.regions = []         # type: List[Dict[str, Any]]
        self._ir_roles = None     # type: Optional[Dict[str, Tuple[str, str]]]

    def ir_roles(self):
        # type: () -> Dict[str, Tuple[str, str]]
        """Cached {text -> (role, field)} map for this slide's IR entry."""
        if self._ir_roles is None:
            self._ir_roles = _ir_role_map(self.ir, self.archetype)
        return self._ir_roles

    def where(self, suffix=""):
        # type: (str) -> str
        base = "slide %d" % self.index
        if self.archetype:
            base += " (%s)" % self.archetype
        return "%s / %s" % (base, suffix) if suffix else base


def slide_background_hex(slide, theme):
    # type: (Any, ThemeColors) -> str
    """Resolve a slide's paint colour, falling back through layout, master, white."""
    sources = []
    sources.append(slide)
    try:
        sources.append(slide.slide_layout)
        sources.append(slide.slide_layout.slide_master)
    except Exception:
        pass
    for src in sources:
        try:
            c_sld = src._element.find(qn("p:cSld"))
            if c_sld is None:
                continue
            bg = c_sld.find(qn("p:bg"))
            if bg is None:
                continue
            bg_pr = bg.find(qn("p:bgPr"))
            if bg_pr is not None:
                info = fill_info(bg_pr, theme)
                if info["kind"] in ("solid", "gradient") and info["colors"]:
                    return info["colors"][0]
            bg_ref = bg.find(qn("p:bgRef"))
            if bg_ref is not None:
                hexed = _first_color_child(bg_ref, theme)
                if hexed:
                    return hexed
        except Exception:
            continue
    return "#FFFFFF"


# ---------------------------------------------------------------------------
# ROLE INFERENCE
# ---------------------------------------------------------------------------

TITLE_ROLES = ("cover", "section", "title")

#: When two type-scale roles share a point size, prefer the earlier one.
ROLE_PREFERENCE = ("title", "body", "subtitle", "cardTitle", "bodySmall", "label",
                   "caption", "eyebrow", "footer", "section", "cover", "statNumber")

#: IR field -> the type-scale role its text is rendered with.
IR_FIELD_ROLE = {
    "eyebrow": "eyebrow",
    "title": "title",
    "subtitle": "subtitle",
    "deck": "subtitle",
    "intro": "subtitle",
    "kicker": "body",
    "body": "body",
    "quote": "section",
    "attrib": "label",
    "attribSub": "caption",
    "label": "label",
    "meta": "caption",
    "cta": "subtitle",
    "contact": "body",
    "caption": "caption",
}


def _ir_role_map(ir_slide, archetype):
    # type: (Optional[Dict[str, Any]], Optional[str]) -> Dict[str, Tuple[str, str]]
    """{normalised text -> (type-scale role, IR field name)} for one IR slide."""
    out = {}  # type: Dict[str, Tuple[str, str]]
    if not isinstance(ir_slide, dict):
        return out
    for field, role in IR_FIELD_ROLE.items():
        value = ir_slide.get(field)
        if isinstance(value, str) and value.strip():
            if field == "title":
                if archetype in ("cover", "closing"):
                    role = "cover"
                elif archetype == "section-break":
                    role = "section"
            out[_norm_text(value)] = (role, field)
    for key in ("items", "stats", "cells", "rows", "columns", "steps", "stages", "photos"):
        seq = ir_slide.get(key)
        if not isinstance(seq, list):
            continue
        for item in seq:
            if isinstance(item, str) and item.strip():
                out.setdefault(_norm_text(item), ("label", "item"))
            elif isinstance(item, dict):
                for field, role in (("subtitle", "cardTitle"), ("label", "cardTitle"),
                                    ("body", "bodySmall"), ("value", "statNumber"),
                                    ("caption", "caption"), ("number", "cardTitle")):
                    value = item.get(field)
                    if isinstance(value, str) and value.strip():
                        out.setdefault(_norm_text(value), (role, field))
    return out


def _norm_text(s):
    # type: (Any) -> str
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def infer_role(rec, ctx, brand):
    # type: (ShapeRec, SlideCtx, Dict[str, Any]) -> Tuple[Optional[str], Optional[str]]
    """(type-scale role, semantic name) for a text shape.

    The semantic name is the IR field or grammar region the text belongs to
    ("title", "quote", "eyebrow", ...). It is what distinguishes a pull quote
    from a slide title: both are set in the "section" type role.
    Resolution order: IR text match, then geometry, then point size.
    """
    text_key = _norm_text(rec.text)
    if text_key:
        ir_map = ctx.ir_roles()
        if text_key in ir_map:
            return ir_map[text_key]

    best = None
    best_d = None
    for region in ctx.regions:
        if region["kind"] != "text" or not region["role"]:
            continue
        d = abs(rec.left - region["x"]) + abs(rec.top - region["y"])
        if d <= REGION_TOL_IN * 2 and (best_d is None or d < best_d):
            best_d = d
            best = (region["role"], region["name"])
    if best:
        return best

    sizes = [r["size"] for para in rec.runs for r in para["runs"] if r["size"]]
    if sizes:
        target = float(max(sizes))
        scale = ((brand.get("type") or {}).get("deckScalePt") or {})
        candidates = []
        for role, spec in scale.items():
            try:
                delta = abs(float(spec.get("size", 0)) - target)
            except (TypeError, ValueError):
                continue
            if delta <= SCALE_TOL_PT:
                rank = ROLE_PREFERENCE.index(role) if role in ROLE_PREFERENCE \
                    else len(ROLE_PREFERENCE)
                candidates.append((delta, rank, role))
        if candidates:
            candidates.sort()
            return (candidates[0][2], None)
    return (None, None)


def role_is_eyebrow(rec, ctx, grammar):
    # type: (ShapeRec, SlideCtx, Grammar) -> bool
    if rec.semantic == "eyebrow" or rec.role == "eyebrow":
        return True
    header_y = float((grammar.chrome.get("headerRule") or {}).get("y", 0.961))
    sizes = [r["size"] for para in rec.runs for r in para["runs"] if r["size"]]
    if rec.top + rec.height <= header_y + 0.05 and (not sizes or max(sizes) <= 12.0):
        return True
    return False


# ---------------------------------------------------------------------------
# LOGO REGISTRY
# ---------------------------------------------------------------------------

class LogoRegistry(object):
    """Known logo variants for a brand, indexed by file hash and aspect."""

    def __init__(self, brand):
        # type: (Dict[str, Any]) -> None
        self.by_sha1 = {}   # type: Dict[str, str]
        self.aspects = {}   # type: Dict[str, float]
        self.files = {}     # type: Dict[str, str]
        logo = brand.get("logo") or {}
        base = brand.get("_dir") or bl.brand_dir(brand.get("id", ""))
        for name, spec in (logo.get("variants") or {}).items():
            if not isinstance(spec, dict):
                continue
            aspect = spec.get("aspect")
            path = os.path.join(base, str(spec.get("file") or ""))
            if spec.get("file") and os.path.isfile(path):
                self.files[name] = path
                try:
                    with open(path, "rb") as fh:
                        self.by_sha1[hashlib.sha1(fh.read()).hexdigest()] = name
                except IOError:
                    pass
                if aspect is None:
                    aspect = _image_aspect(path)
            if aspect:
                try:
                    self.aspects[name] = float(aspect)
                except (TypeError, ValueError):
                    pass

    def identify(self, rec):
        # type: (ShapeRec) -> Tuple[Optional[str], Optional[str]]
        """(variant name, how it was matched) for a picture shape."""
        if not rec.is_picture:
            return (None, None)
        sha1 = None
        filename = ""
        native = None  # type: Optional[float]
        try:
            image = rec.shape.image
            sha1 = image.sha1
            filename = str(image.filename or "")
            size = image.size
            if size and size[1]:
                native = float(size[0]) / float(size[1])
        except Exception:
            pass

        if sha1 and sha1 in self.by_sha1:
            return (self.by_sha1[sha1], "hash")

        haystack = " ".join([filename, rec.name, _shape_descr(rec.el)]).lower()
        if "logo" in haystack or "lockup" in haystack or "wordmark" in haystack:
            return (self._nearest_aspect(native), "name")

        if native is not None:
            match = self._nearest_aspect(native, tol=0.03)
            if match:
                return (match, "aspect")
        return (None, None)

    def _nearest_aspect(self, native, tol=None):
        # type: (Optional[float], Optional[float]) -> Optional[str]
        if native is None or not self.aspects:
            return None
        best, best_d = None, None
        for name, aspect in self.aspects.items():
            if aspect <= 0:
                continue
            d = abs(native - aspect) / aspect
            if best_d is None or d < best_d:
                best, best_d = name, d
        if best_d is None:
            return None
        if tol is not None and best_d > tol:
            return None
        return best


def _image_aspect(path):
    # type: (str) -> Optional[float]
    try:
        from PIL import Image
        with Image.open(path) as img:
            w, h = img.size
        if h:
            return float(w) / float(h)
    except Exception:
        return None
    return None


# ---------------------------------------------------------------------------
# THE VALIDATOR
# ---------------------------------------------------------------------------

class DeckValidator(object):

    def __init__(self, path, brand, grammar, ir=None):
        # type: (str, Dict[str, Any], Grammar, Optional[Dict[str, Any]]) -> None
        self.path = path
        self.brand = brand
        self.grammar = grammar
        self.ir = ir or None
        self.report = bl.Report(target=path, brand=brand.get("id", ""), kind=KIND)
        self.palette = bl.palette_index(brand)
        self.superseded = bl.superseded_map(brand)
        self.forbidden_text = bl.forbidden_text_colors(brand)
        self.forbidden_phrases = ((brand.get("voice") or {}).get("forbiddenPhrases") or [])
        self.forbidden_chars = ((brand.get("voice") or {}).get("forbiddenChars") or ["!"])
        self.families = self._approved_families()
        self.min_body_pt = float((brand.get("type") or {}).get("minBodyPt", 10.5))
        self.scale_sizes = self._scale_sizes()
        self.neg_tracking_above = float(
            (brand.get("type") or {}).get("negativeTrackingAbovePt", 22))
        self.logos = LogoRegistry(brand)
        self._seen = set()  # type: set
        self.contexts = []  # type: List[SlideCtx]

    # -- plumbing -----------------------------------------------------------

    def _approved_families(self):
        # type: () -> Dict[str, str]
        tmap = ((self.brand.get("type") or {}).get("weightToPptxFamily") or {})
        out = {}
        for _weight, family in tmap.items():
            out[_norm_family(family)] = str(family)
        base = (self.brand.get("type") or {}).get("family")
        if base:
            out.setdefault(_norm_family(base), str(base))
        return out

    def _scale_sizes(self):
        # type: () -> List[float]
        sizes = []
        for spec in ((self.brand.get("type") or {}).get("deckScalePt") or {}).values():
            try:
                sizes.append(float(spec.get("size")))
            except (TypeError, ValueError):
                continue
        return sorted(set(sizes))

    def add(self, vid, severity, where, found="", expected="", rule="", fix="", key=None):
        # type: (str, str, str, str, str, str, str, Any) -> None
        """Add a violation, suppressing exact repeats of the same (key)."""
        dedupe = key if key is not None else (vid, where, found, expected)
        if dedupe in self._seen:
            return
        self._seen.add(dedupe)
        self.report.add(vid, severity=severity, where=where, found=found,
                        expected=expected, rule=rule, fix=fix)

    # -- entry point --------------------------------------------------------

    def run(self):
        # type: () -> bl.Report
        prs = Presentation(self.path)
        theme_cache = {}  # type: Dict[Any, ThemeColors]
        slides = list(prs.slides)
        ir_slides = (self.ir or {}).get("slides") or []

        ir_brand = str((self.ir or {}).get("brand") or "")
        if ir_brand and ir_brand != self.brand.get("id"):
            self.add("STRUCTURE.IR_MISMATCH", "info", "deck",
                     found="the IR names brand %r" % ir_brand,
                     expected="validation is running against %r" % self.brand.get("id"),
                     rule="A deck is validated against the brand that authored it.",
                     fix="Re-run with --brand %s, or fix the IR's brand field." % ir_brand)

        if self.ir and len(ir_slides) != len(slides):
            self.add("STRUCTURE.IR_MISMATCH", "info", "deck",
                     found="%d slides in the pptx, %d in the IR" % (len(slides), len(ir_slides)),
                     expected="the same number of slides in both",
                     rule="The IR and the rendered deck must describe the same deck.",
                     fix="Rebuild the deck from the IR you are validating against.")

        for i, slide in enumerate(slides, start=1):
            ctx = SlideCtx()
            ctx.index = i
            ctx.slide = slide
            ctx.theme = _theme_for_slide(slide, theme_cache)
            ctx.bg_hex = slide_background_hex(slide, ctx.theme)
            ctx.ir = ir_slides[i - 1] if i - 1 < len(ir_slides) else None

            recs = []  # type: List[ShapeRec]
            flatten_shapes(slide.shapes, _identity_xform, recs, [0], "")
            ctx.recs = recs

            for rec in recs:
                sp_pr = rec.el.find(qn("p:spPr"))
                rec.fill = fill_info(sp_pr, ctx.theme)
                ln = sp_pr.find(qn("a:ln")) if sp_pr is not None else None
                rec.line = fill_info(ln, ctx.theme)
                if getattr(rec.shape, "has_text_frame", False):
                    rec.runs = extract_paragraphs(rec.shape.text_frame, ctx.theme)
                    rec.text = "\n".join(p["text"] for p in rec.runs).strip()
                if rec.is_picture:
                    rec.logo_variant, rec.logo_conf = self.logos.identify(rec)

            if isinstance(ctx.ir, dict) and ctx.ir.get("archetype"):
                ctx.archetype = str(ctx.ir["archetype"])
                ctx.archetype_source = "ir"
            else:
                ctx.archetype = infer_archetype(recs, self.grammar, i, len(slides))
                ctx.archetype_source = "inferred" if ctx.archetype else "none"
            ctx.regions = self.grammar.regions(ctx.archetype)

            for rec in recs:
                if rec.text:
                    rec.role, rec.semantic = infer_role(rec, ctx, self.brand)

            self.contexts.append(ctx)
            self.check_slide(ctx)

        self.check_structure(len(slides))
        self.check_learned_rules()
        return self.report

    # -- per slide ----------------------------------------------------------

    def check_slide(self, ctx):
        # type: (SlideCtx) -> None
        gradient_shapes = []  # type: List[ShapeRec]
        has_content = False

        for rec in ctx.recs:
            where = ctx.where(rec.path)
            if rec.fill.get("kind") == "gradient":
                gradient_shapes.append(rec)
            if rec.text or rec.is_picture or rec.is_chart or rec.is_table:
                has_content = True

            self.check_shape_colors(ctx, rec, where)
            self.check_geometry(ctx, rec, where)

            if rec.runs:
                self.check_text_frame(ctx, rec, where)
            if rec.is_table:
                self.check_table(ctx, rec, where)
            if rec.is_chart:
                self.check_chart(ctx, rec, where)
            if rec.is_picture:
                self.check_picture(ctx, rec, where)

        if len(gradient_shapes) > 1:
            self.add("COLOR.TWO_GRADIENTS", "warn", ctx.where(),
                     found="%d gradient-filled shapes (%s)"
                           % (len(gradient_shapes),
                              ", ".join(r.path for r in gradient_shapes[:4])),
                     expected="at most %s gradient per surface"
                              % ((self.brand.get("colorRules") or {}).get(
                                  "maxGradientsPerSurface", 1)),
                     rule="One gradient per surface. Competing gradients read as noise.",
                     fix="Keep the hero gradient and flatten the others to a solid brand colour.")

        if not has_content:
            self.add("CONTENT.EMPTY_SLIDE", "warn", ctx.where(),
                     found="no text, picture, chart or table",
                     expected="every slide carries content",
                     rule="An empty slide is a build error, not a design choice.",
                     fix="Populate the slide from the IR, or delete it.")

        self.check_logo(ctx)
        self.check_overlap(ctx)
        self.check_contrast(ctx)

    # -- colour -------------------------------------------------------------

    def judge_color(self, hexed, is_text):
        # type: (str, bool) -> Optional[Tuple[str, str, str, str, str]]
        """(id, severity, found, expected, fix) for a colour, or None when approved."""
        try:
            value = bl.normalize_hex(hexed)
        except ValueError:
            return None

        if value in self.superseded:
            canonical = self.superseded[value]
            token = self.palette.get(canonical, "the canonical token")
            return ("COLOR.SUPERSEDED", "error", value,
                    "%s (%s)" % (canonical, token),
                    "Recolour to %s. The template palette was superseded by the design "
                    "system on the brand's ruling date." % canonical)

        if is_text:
            if bl.delta_e(value, "#000000") <= BLACK_DELTA_E:
                default_text = ((self.brand.get("colorRules") or {}).get("defaultText")
                                or "#0F0A6C")
                return ("COLOR.BLACK_TEXT", "error", value, default_text,
                        "Never set text in pure or near black. Use the brand's default "
                        "text colour %s." % default_text)
            if value in self.forbidden_text:
                reason = self.forbidden_text[value]
                return ("COLOR.FORBIDDEN_TEXT", "error", value,
                        "a text colour that is not %s" % value, reason)

        if value in self.palette:
            return None

        name, near_hex, dist = bl.nearest_token(value, self.palette)
        if near_hex is None:
            return ("COLOR.OFF_PALETTE", "error", value, "an approved brand colour",
                    "The brand profile declares no palette to snap to.")
        return ("COLOR.OFF_PALETTE", "error", value,
                "%s (%s)" % (near_hex, name),
                "Snap to %s (%s); delta E %.1f from what was used." % (near_hex, name, dist))

    def report_color(self, vid_where, hexed, is_text, context_label):
        # type: (str, str, bool, str) -> None
        verdict = self.judge_color(hexed, is_text)
        if verdict is None:
            return
        vid, severity, found, expected, fix = verdict
        rule = {
            "COLOR.SUPERSEDED": "brand.color.superseded lists template colours that are "
                                "not approved.",
            "COLOR.BLACK_TEXT": "Pure black is never used for text.",
            "COLOR.FORBIDDEN_TEXT": "brand.colorRules.forbiddenText bans this colour as text.",
            "COLOR.OFF_PALETTE": "Every colour must come from the brand palette.",
        }[vid]
        self.add(vid, severity, vid_where,
                 found="%s (%s)" % (found, context_label),
                 expected=expected, rule=rule, fix=fix,
                 key=(vid, vid_where, found, context_label))

    def check_shape_colors(self, ctx, rec, where):
        # type: (SlideCtx, ShapeRec, str) -> None
        for kind_label, info in (("fill", rec.fill), ("line", rec.line)):
            kind = info.get("kind")
            if kind in ("solid", "gradient", "pattern"):
                for hexed in info.get("colors") or []:
                    self.report_color(where, hexed, False, kind_label)

    def check_text_frame(self, ctx, rec, where):
        # type: (SlideCtx, ShapeRec, str) -> None
        self.check_runs(rec.runs, where, rec.role)
        self.check_copy(ctx, rec, where)
        self.check_overflow(ctx, rec, where)

    def check_runs(self, paras, where, role):
        # type: (List[Dict[str, Any]], str, Optional[str]) -> None
        expected_family = self._expected_family(role)
        for para in paras:
            for run in para["runs"]:
                if not (run["text"] or "").strip():
                    continue
                loc = "%s / p%d r%d" % (where, para["index"] + 1, run["index"] + 1)

                family = run["family"]
                if family is None:
                    self.add("TYPE.INHERITED_FONT", "error", loc,
                             found="no explicit font family on the run",
                             expected=expected_family,
                             rule="An unset font family silently renders as the Office "
                                  "theme font.",
                             fix="Set the run font to %s." % expected_family,
                             key=("TYPE.INHERITED_FONT", where, para["index"], run["index"]))
                elif _norm_family(family) not in self.families:
                    self.add("TYPE.OFF_FAMILY", "error", loc,
                             found=str(family),
                             expected=", ".join(sorted(self.families.values())),
                             rule="Only the brand type families may ship in a deck.",
                             fix="Replace %s with %s." % (family, expected_family),
                             key=("TYPE.OFF_FAMILY", where, _norm_family(family)))

                if run["bold"]:
                    self.add("TYPE.SYNTHETIC_BOLD", "error", loc,
                             found="bold=True on %s" % (family or "an inherited family"),
                             expected="bold=False with the weight carried by the family name",
                             rule="Weights are family names in this system; synthetic bold "
                                  "distorts the letterforms.",
                             fix="Set bold=False and use %s." % self._heaviest_family(),
                             key=("TYPE.SYNTHETIC_BOLD", where, para["index"], run["index"]))

                if run["italic"]:
                    self.add("TYPE.ITALIC", "error", loc,
                             found="italic=True",
                             expected="italic=False",
                             rule="brand.type.italicsAllowed is false.",
                             fix="Remove the italic and use a heavier family for emphasis.",
                             key=("TYPE.ITALIC", where, para["index"], run["index"]))

                size = run["size"]
                if size is not None:
                    if float(size) < self.min_body_pt - 1e-6:
                        self.add("TYPE.BELOW_MIN", "error", loc,
                                 found="%.4gpt" % float(size),
                                 expected=">= %.4gpt" % self.min_body_pt,
                                 rule="brand.type.minBodyPt is the floor for shipped text.",
                                 fix="Raise the size to at least %.4gpt, or cut copy."
                                     % self.min_body_pt,
                                 key=("TYPE.BELOW_MIN", where, round(float(size), 2)))
                    elif self.scale_sizes and not any(
                            abs(float(size) - s) <= SCALE_TOL_PT for s in self.scale_sizes):
                        self.add("TYPE.OFF_SCALE", "warn", loc,
                                 found="%.4gpt" % float(size),
                                 expected="one of %s"
                                          % ", ".join("%.4g" % s for s in self.scale_sizes),
                                 rule="Sizes come from brand.type.deckScalePt.",
                                 fix="Snap to the nearest scale step.",
                                 key=("TYPE.OFF_SCALE", where, round(float(size), 2)))

                spc = run["spc"]
                if spc is not None and spc < 0:
                    effective = float(size) if size is not None else 0.0
                    if effective < self.neg_tracking_above:
                        self.add("TYPE.TRACKING", "warn", loc,
                                 found="tracking %.2fpt at %s"
                                       % (spc / 100.0,
                                          ("%.4gpt" % effective) if size else "an unset size"),
                                 expected="negative tracking only above %.4gpt"
                                          % self.neg_tracking_above,
                                 rule="brand.type.negativeTrackingAbovePt guards small text "
                                      "from being tightened.",
                                 fix="Set tracking to 0 at this size.",
                                 key=("TYPE.TRACKING", where, para["index"], run["index"]))

                fill = run["fill"]
                if fill.get("kind") == "gradient":
                    self.add("COLOR.GRADIENT_TEXT", "error", loc,
                             found="gradient fill on a text run",
                             expected="a solid brand colour",
                             rule="brand.colorRules.gradientTextForbidden is true.",
                             fix="Set the run to a solid palette colour.",
                             key=("COLOR.GRADIENT_TEXT", where, para["index"], run["index"]))
                elif fill.get("kind") in ("solid", "pattern"):
                    for hexed in fill.get("colors") or []:
                        self.report_color(loc, hexed, True, "text")

    def _heaviest_family(self):
        # type: () -> str
        tmap = ((self.brand.get("type") or {}).get("weightToPptxFamily") or {})
        best = None
        for weight, family in tmap.items():
            try:
                w = int(str(weight))
            except ValueError:
                continue
            if best is None or w > best[0]:
                best = (w, str(family))
        if best:
            return best[1]
        return str((self.brand.get("type") or {}).get("family") or "the brand family")

    def _expected_family(self, role):
        # type: (Optional[str]) -> str
        weight = 400
        if role:
            try:
                weight = int((bl.type_role(self.brand, role) or {}).get("weight", 400))
            except KeyError:
                weight = 400
        try:
            return bl.weight_to_family(self.brand, weight)
        except ValueError:
            return str((self.brand.get("type") or {}).get("family") or "the brand family")

    # -- copy / voice -------------------------------------------------------

    def check_copy(self, ctx, rec, where):
        # type: (SlideCtx, ShapeRec, str) -> None
        text = rec.text
        if not text:
            return

        for phrase in bl.find_forbidden_phrases(text, self.forbidden_phrases):
            self.add("CONTENT.PLACEHOLDER", "error", where,
                     found=phrase,
                     expected="real client copy",
                     rule="brand.voice.forbiddenPhrases must never ship.",
                     fix="Replace %r with the real sentence for this slide." % phrase,
                     key=("CONTENT.PLACEHOLDER", where, phrase))

        for char in self.forbidden_chars:
            if char and char in text:
                if char == "!":
                    vid = "VOICE.EXCLAMATION"
                    rule = "The brand voice uses no exclamation marks."
                    fix = "Delete the exclamation mark and let the sentence carry the weight."
                else:
                    vid = "VOICE.EXCLAMATION"
                    rule = "brand.voice.forbiddenChars bans this character."
                    fix = "Remove %r from the copy." % char
                self.add(vid, "error", where,
                         found="%r in %r" % (char, _snippet(text)),
                         expected="copy without %r" % char,
                         rule=rule, fix=fix,
                         key=(vid, where, char))

        is_eyebrow = role_is_eyebrow(rec, ctx, self.grammar)
        paras = [p for p in rec.runs if (p["text"] or "").strip()]

        if not is_eyebrow:
            for para in paras:
                body = para["text"].strip()
                if not bl.is_sentence_case(body):
                    self.add("VOICE.TITLE_CASE", "warn", where,
                             found=_snippet(body),
                             expected="sentence case",
                             rule="brand.voice.case is sentence; only the eyebrow is upper.",
                             fix="Lower-case everything after the first word except proper "
                                 "nouns and acronyms.",
                             key=("VOICE.TITLE_CASE", where, _norm_text(body)))

        max_bullets = int(self.grammar.rule("maxBulletsPerSlide", 6))
        if len(paras) > max_bullets:
            self.add("CONTENT.BULLET_COUNT", "warn", where,
                     found="%d paragraphs" % len(paras),
                     expected="at most %d" % max_bullets,
                     rule="grammar.deckRules.maxBulletsPerSlide caps one body block.",
                     fix="Split the slide, or promote the extra points to their own slide.",
                     key=("CONTENT.BULLET_COUNT", where))

        max_words = int(self.grammar.rule("maxWordsPerBullet", 18))
        counts = [bl.word_count(p["text"]) for p in paras]
        list_like = any(p["bulleted"] for p in paras) or (
            len(paras) >= 2 and min(counts) <= max_words)
        if list_like:
            for para, words in zip(paras, counts):
                if words > max_words:
                    self.add("CONTENT.LONG_BULLET", "warn", where,
                             found="%d words: %s" % (words, _snippet(para["text"])),
                             expected="at most %d words per bullet" % max_words,
                             rule="grammar.deckRules.maxWordsPerBullet keeps list items "
                                  "scannable.",
                             fix="Cut the bullet to one clause, or move the detail to the "
                                 "speaker notes.",
                             key=("CONTENT.LONG_BULLET", where, _norm_text(para["text"])))

        max_title = int(self.grammar.rule("maxTitleChars", 70))
        is_title = rec.semantic == "title" or (rec.semantic is None
                                               and rec.role in TITLE_ROLES)
        if is_title and len(text) > max_title:
            self.add("CONTENT.LONG_TITLE", "warn", where,
                     found="%d characters: %s" % (len(text), _snippet(text)),
                     expected="at most %d characters" % max_title,
                     rule="grammar.deckRules.maxTitleChars caps a slide title.",
                     fix="Shorten the title; move the qualifier into the deck line.",
                     key=("CONTENT.LONG_TITLE", where))

    # -- layout -------------------------------------------------------------

    def check_geometry(self, ctx, rec, where):
        # type: (SlideCtx, ShapeRec, str) -> None
        cw, ch = self.grammar.canvas_w, self.grammar.canvas_h
        right = rec.left + rec.width
        bottom = rec.top + rec.height
        off = []
        if rec.left < -CANVAS_TOL_IN:
            off.append("left %.3fin" % rec.left)
        if rec.top < -CANVAS_TOL_IN:
            off.append("top %.3fin" % rec.top)
        if right > cw + CANVAS_TOL_IN:
            off.append("right %.3fin" % right)
        if bottom > ch + CANVAS_TOL_IN:
            off.append("bottom %.3fin" % bottom)
        if off:
            self.add("LAYOUT.OFF_CANVAS", "error", where,
                     found=", ".join(off),
                     expected="0..%.3f x 0..%.3f in" % (cw, ch),
                     rule="Nothing may hang off the %s canvas."
                          % ((self.grammar.data.get("canvas") or {}).get("ratio", "16:9")),
                     fix="Move or resize the shape back inside the canvas.",
                     key=("LAYOUT.OFF_CANVAS", where))

        if rec.width <= 0 or rec.height <= 0:
            return
        allowed = [0.0,
                   float(self.grammar.grid.get("contentLeft", 0.869)),
                   float(self.grammar.grid.get("bleedLeft", 0.41))]
        allowed.extend(self.grammar.column_xs())
        # A shape sitting exactly on one of its archetype's declared region x values,
        # or on an item x inside one of its sets, is on-grid by definition (hero panels
        # and stat columns are deliberately off the 12-column rhythm).
        allowed.extend(region["x"] for region in ctx.regions)
        allowed.extend(self.grammar.set_xs(ctx.archetype))
        # Item content is inset from its container by the archetype's own dx values.
        dxs = self.grammar.spec_dxs(ctx.archetype)
        if dxs:
            allowed.extend(base + dx for base in list(allowed) for dx in dxs)
        if not any(abs(rec.left - x) <= 0.02 for x in allowed):
            self.add("LAYOUT.OFF_GRID", "info", where,
                     found="left %.3fin" % rec.left,
                     expected="0, %.3f (bleed), %.3f (content) or a column x"
                              % (float(self.grammar.grid.get("bleedLeft", 0.41)),
                                 float(self.grammar.grid.get("contentLeft", 0.869))),
                     rule="Left edges sit on the 12-column grid.",
                     fix="Snap the left edge to the nearest column x.",
                     key=("LAYOUT.OFF_GRID", where))

    def check_overflow(self, ctx, rec, where):
        # type: (SlideCtx, ShapeRec, str) -> None
        if not rec.text:
            return
        try:
            tf = rec.shape.text_frame
        except Exception:
            return

        ml = bl.emu_to_in(tf.margin_left or 0)
        mr = bl.emu_to_in(tf.margin_right or 0)
        mt = bl.emu_to_in(tf.margin_top or 0)
        mb = bl.emu_to_in(tf.margin_bottom or 0)
        box_w = rec.width - ml - mr
        box_h = rec.height - mt - mb
        if box_w <= 0.05 or box_h <= 0.05:
            return

        wrap = tf.word_wrap
        default_size = self._role_size(rec.role)
        pct = float(self.grammar.rule("overflowTolerancePct", 2.0))
        tol = pct / 100.0

        ratio = self._leading_ratio(rec.role)
        needed_h = 0.0
        needed_w = 0.0
        for para in rec.runs:
            text = para["text"]
            size = paragraph_size(para, default_size)
            family = paragraph_family(para, None)
            line_h = size * ratio / 72.0
            if not text.strip():
                needed_h += line_h
                continue
            if wrap is False:
                # Unwrapped text keeps every paragraph on one line and spills sideways.
                lines = text.split("\n")
                needed_h += len(lines) * line_h
                char_w = bl.GLYPH_WIDTH_EM.get(
                    _norm_family(family), bl.DEFAULT_GLYPH_WIDTH_EM) * size / 72.0
                needed_w = max(needed_w, max(len(line) for line in lines) * char_w)
            else:
                # brandlib does the wrapping; the brand's own leading does the spacing.
                _fits, height = bl.fits_in_box(text, size, box_w, box_h, family)
                lines = max(1, int(round(height / (bl.LEADING_FACTOR * size / 72.0))))
                needed_h += lines * line_h

        if needed_h > box_h * (1.0 + tol) + 1e-6:
            self.add("LAYOUT.OVERFLOW", "error", where,
                     found="text needs %.3fin of height in a %.3fin box (%.1f%% over)"
                           % (needed_h, box_h,
                              (needed_h / box_h - 1.0) * 100.0 if box_h else 0.0),
                     expected="<= %.3fin (%.4g%% tolerance)" % (box_h * (1.0 + tol), pct),
                     rule="Estimated text height must fit its shape.",
                     fix="Cut copy, drop one step on the type scale, or move the block to "
                         "its own slide.",
                     key=("LAYOUT.OVERFLOW", where, "h"))
        elif needed_w > box_w * (1.0 + tol) + 1e-6:
            self.add("LAYOUT.OVERFLOW", "error", where,
                     found="unwrapped text needs %.3fin of width in a %.3fin box"
                           % (needed_w, box_w),
                     expected="<= %.3fin (%.4g%% tolerance)" % (box_w * (1.0 + tol), pct),
                     rule="A text frame with word wrap off must still fit its shape.",
                     fix="Turn word wrap on for this text frame, or shorten the line.",
                     key=("LAYOUT.OVERFLOW", where, "w"))

    def _leading_ratio(self, role):
        # type: (Optional[str]) -> float
        """Line height as a multiple of point size, from the brand's own type scale."""
        if role:
            try:
                spec = bl.type_role(self.brand, role) or {}
                size = float(spec.get("size") or 0.0)
                leading = float(spec.get("leading") or 0.0)
                if size > 0 and leading > 0:
                    return max(1.0, leading / size)
            except (KeyError, TypeError, ValueError):
                pass
        return bl.LEADING_FACTOR

    def _role_size(self, role):
        # type: (Optional[str]) -> float
        if role:
            try:
                return float((bl.type_role(self.brand, role) or {}).get("size", 12.0))
            except KeyError:
                pass
        try:
            return float(bl.type_role(self.brand, "body").get("size", 12.0))
        except KeyError:
            return 12.0

    def check_overlap(self, ctx):
        # type: (SlideCtx) -> None
        tol = float(self.grammar.rule("overlapToleranceIn", 0.02))
        limit = tol * tol
        bearing = [r for r in ctx.recs if r.text and r.width > 0 and r.height > 0]
        for i in range(len(bearing)):
            for j in range(i + 1, len(bearing)):
                a, b = bearing[i], bearing[j]
                area = bl.rect_overlap(a.rect, b.rect)
                if area > limit:
                    self.add("LAYOUT.OVERLAP", "warn", ctx.where("%s + %s" % (a.path, b.path)),
                             found="%.4f sq in of overlap" % area,
                             expected="<= %.4f sq in" % limit,
                             rule="grammar.deckRules.overlapToleranceIn: text blocks do not "
                                  "collide.",
                             fix="Move one block onto its own grid column or row.",
                             key=("LAYOUT.OVERLAP", ctx.index, a.path, b.path))

    # -- accessibility ------------------------------------------------------

    def effective_background(self, ctx, rec):
        # type: (SlideCtx, ShapeRec) -> List[str]
        """Colours painted behind a text shape, nearest enclosing surface first."""
        if rec.area <= 0:
            return [ctx.bg_hex]
        best = None  # type: Optional[ShapeRec]
        for other in ctx.recs:
            if other is rec or other.z >= rec.z:
                continue
            if other.fill.get("kind") not in ("solid", "gradient"):
                continue
            colors = [c for c in (other.fill.get("colors") or []) if c]
            if not colors:
                continue
            covered = bl.rect_overlap(other.rect, rec.rect)
            if covered / rec.area >= CONTAINMENT_RATIO:
                best = other
        if best is not None:
            return [c for c in (best.fill.get("colors") or []) if c] or [ctx.bg_hex]
        return [ctx.bg_hex]

    def check_contrast(self, ctx):
        # type: (SlideCtx) -> None
        for rec in ctx.recs:
            if not rec.runs:
                continue
            backgrounds = self.effective_background(ctx, rec)
            where = ctx.where(rec.path)
            default_size = self._role_size(rec.role)
            for para in rec.runs:
                for run in para["runs"]:
                    if not (run["text"] or "").strip():
                        continue
                    fill = run["fill"]
                    if fill.get("kind") != "solid" or not fill.get("colors"):
                        continue
                    fg = fill["colors"][0]
                    if not fg:
                        continue
                    size = float(run["size"]) if run["size"] else default_size
                    worst = None
                    for bg in backgrounds:
                        try:
                            ok, required, actual = bl.passes_contrast(fg, bg, size, run["bold"])
                        except ValueError:
                            continue
                        if worst is None or actual < worst[2]:
                            worst = (ok, required, actual, bg)
                    if worst is None or worst[0]:
                        continue
                    _ok, required, actual, bg = worst
                    mandated = bl.on_surface_text(self.brand, bg)
                    expected = ("%s on %s" % (mandated, bg)) if mandated \
                        else "a text colour reaching %.1f:1 on %s" % (required, bg)
                    self.add("A11Y.CONTRAST", "error",
                             "%s / p%d r%d" % (where, para["index"] + 1, run["index"] + 1),
                             found="%s on %s is %.2f:1 at %.4gpt" % (fg, bg, actual, size),
                             expected=expected,
                             rule="WCAG AA needs %.1f:1 at this size." % required,
                             fix=("Use %s on this surface." % mandated) if mandated
                                 else "Darken the text or lighten the surface until it "
                                      "reaches %.1f:1." % required,
                             key=("A11Y.CONTRAST", ctx.index, rec.path, fg, bg,
                                  round(size, 1)))

    def check_picture(self, ctx, rec, where):
        # type: (SlideCtx, ShapeRec, str) -> None
        descr = _shape_descr(rec.el).strip()
        meaningful = bool(descr) and not _FILENAMEISH_RE.match(descr) \
            and not _IMAGE_EXT_RE.search(descr)
        if not meaningful:
            self.add("A11Y.ALT_TEXT", "info", where,
                     found=("descr=%r" % descr) if descr else "no descr attribute",
                     expected="a one-line description of what the image shows",
                     rule="Every picture carries alt text.",
                     fix="Set the picture's alt text (cNvPr@descr) to a short description.",
                     key=("A11Y.ALT_TEXT", ctx.index, rec.path))

    # -- logo ---------------------------------------------------------------

    def check_logo(self, ctx):
        # type: (SlideCtx) -> None
        logo_cfg = self.brand.get("logo") or {}
        pictures = [r for r in ctx.recs if r.is_picture]
        logos = [r for r in pictures if r.logo_variant]

        needs_logo = ctx.archetype in ("cover", "closing")
        if ctx.archetype is None and ctx.index == 1:
            needs_logo = True
        if needs_logo and self.grammar.rule("requireLogoOnCoverAndClosing", True):
            if not pictures:
                self.add("LOGO.MISSING", "error", ctx.where(),
                         found="no image on the slide",
                         expected="the brand logo",
                         rule="grammar.deckRules.requireLogoOnCoverAndClosing.",
                         fix="Place the %s logo variant on this slide."
                             % ((logo_cfg.get("variantForBackground") or {}).get("dark",
                                                                                "primary")),
                         key=("LOGO.MISSING", ctx.index))
            elif not logos:
                self.add("LOGO.MISSING", "error", ctx.where(),
                         found="%d picture(s), none matching a brand logo variant"
                               % len(pictures),
                         expected="one of: %s" % (", ".join(sorted(self.logos.aspects.keys()))
                                                  or "a declared logo variant"),
                         rule="grammar.deckRules.requireLogoOnCoverAndClosing.",
                         fix="Place a logo file from brands/%s/assets/logos on this slide."
                             % self.brand.get("id", ""),
                         key=("LOGO.MISSING", ctx.index))

        min_w = float(logo_cfg.get("minWidthIn", 0.0) or 0.0)
        placement = str(logo_cfg.get("placement") or "")
        arch = self.grammar.archetypes.get(ctx.archetype or "")
        chrome_slide = bool(arch and arch.get("chrome"))

        for rec in logos:
            where = ctx.where(rec.path)
            if min_w and rec.width < min_w - 1e-6:
                self.add("LOGO.UNDERSIZE", "error", where,
                         found="%.3fin wide" % rec.width,
                         expected=">= %.3fin" % min_w,
                         rule="brand.logo.minWidthIn is the legibility floor.",
                         fix="Scale the logo to at least %.3fin wide, keeping its aspect."
                             % min_w,
                         key=("LOGO.UNDERSIZE", ctx.index, rec.path))

            declared = self.logos.aspects.get(rec.logo_variant or "")
            if declared and rec.height > 0:
                rendered = rec.width / rec.height
                drift = abs(rendered - declared) / declared
                if drift > LOGO_ASPECT_TOL:
                    self.add("LOGO.DISTORTED", "error", where,
                             found="aspect %.3f (%.3f x %.3f in)"
                                   % (rendered, rec.width, rec.height),
                             expected="aspect %.3f (+/- %.0f%%)"
                                      % (declared, LOGO_ASPECT_TOL * 100),
                             rule="brand.logo.forbidden includes stretching the mark.",
                             fix="Set the height to %.3fin for this width, or the width to "
                                 "%.3fin for this height."
                                 % (rec.width / declared, rec.height * declared),
                             key=("LOGO.DISTORTED", ctx.index, rec.path))

            if chrome_slide and placement == "top-left" \
                    and rec.left > self.grammar.canvas_w / 2.0:
                self.add("LOGO.PLACEMENT", "warn", where,
                         found="left edge at %.3fin" % rec.left,
                         expected="left of %.3fin (brand.logo.placement is top-left)"
                                  % (self.grammar.canvas_w / 2.0),
                         rule="The logo sits top-left on chrome slides.",
                         fix="Move the logo to x=%.3f, y=%.3f."
                             % (float((self.grammar.chrome.get("logo") or {}).get("x", 0.869)),
                                float((self.grammar.chrome.get("logo") or {}).get("y", 0.30))),
                         key=("LOGO.PLACEMENT", ctx.index, rec.path))

            ratio = float(logo_cfg.get("clearSpaceRatio", 1.0) or 1.0)
            pad = rec.height * ratio
            box = (rec.left - pad, rec.top - pad, rec.width + 2 * pad, rec.height + 2 * pad)
            for other in ctx.recs:
                if other is rec or other.area <= 0:
                    continue
                if other.is_picture and other.logo_variant:
                    continue
                box_area = box[2] * box[3]
                if bl.rect_overlap(other.rect, box) <= 0.0:
                    continue
                # A surface that fully contains the clear-space box is the backdrop,
                # not an intrusion.
                if bl.rect_overlap(other.rect, box) >= box_area - 1e-6:
                    continue
                # Nor is a large, textless surface painted behind the logo: that is the
                # panel the mark sits on, not content crowding it.
                if not other.text and other.z < rec.z and other.area >= 4.0 * box_area:
                    continue
                if not other.text and other.fill.get("kind") in (None, "none") \
                        and other.line.get("kind") in (None, "none"):
                    continue
                self.add("LOGO.CLEARSPACE", "warn", ctx.where("%s vs %s" % (rec.path,
                                                                            other.path)),
                         found="%s intrudes into the %.3fin clear space" % (other.path, pad),
                         expected="nothing within %.3fin of the logo" % pad,
                         rule="brand.logo.clearSpaceBasis: %s"
                              % (logo_cfg.get("clearSpaceBasis") or "logo height on all sides"),
                         fix="Move %s outside the clear-space box." % other.path,
                         key=("LOGO.CLEARSPACE", ctx.index, rec.path, other.path))

    # -- tables -------------------------------------------------------------

    def check_table(self, ctx, rec, where):
        # type: (SlideCtx, ShapeRec, str) -> None
        try:
            table = rec.shape.table
        except Exception:
            return
        col_w = [bl.emu_to_in(c.width or 0) for c in table.columns]
        row_h = [bl.emu_to_in(r.height or 0) for r in table.rows]

        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                if getattr(cell, "is_spanned", False):
                    continue
                loc = "%s / cell(%d,%d)" % (where, ri + 1, ci + 1)
                cell_pr = cell._tc.find(qn("a:tcPr"))
                info = fill_info(cell_pr, ctx.theme)
                if info.get("kind") in ("solid", "gradient", "pattern"):
                    for hexed in info.get("colors") or []:
                        self.report_color(loc, hexed, False, "cell fill")

                paras = extract_paragraphs(cell.text_frame, ctx.theme)
                pseudo = ShapeRec()
                pseudo.shape = cell
                pseudo.el = cell._tc
                pseudo.name = "cell(%d,%d)" % (ri + 1, ci + 1)
                pseudo.path = "%s / %s" % (rec.path, pseudo.name)
                pseudo.runs = paras
                pseudo.text = "\n".join(p["text"] for p in paras).strip()
                span_w = int(getattr(cell, "span_width", 1) or 1)
                span_h = int(getattr(cell, "span_height", 1) or 1)
                pseudo.width = sum(col_w[ci:ci + span_w]) if col_w else rec.width
                pseudo.height = sum(row_h[ri:ri + span_h]) if row_h else rec.height
                pseudo.left = rec.left + sum(col_w[:ci])
                pseudo.top = rec.top + sum(row_h[:ri])
                pseudo.role = "bodySmall"
                pseudo.z = rec.z

                self.check_runs(paras, loc, pseudo.role)
                if pseudo.text:
                    self.check_copy(ctx, pseudo, loc)

    # -- charts -------------------------------------------------------------

    def check_chart(self, ctx, rec, where):
        # type: (SlideCtx, ShapeRec, str) -> None
        try:
            blob = rec.shape.chart.part.blob
            root = ET.fromstring(blob)
        except Exception:
            return

        for ser in root.iter("%sser" % _C):
            if list(ser.iter("%sgradFill" % _A)):
                self.add("COLOR.CHART_GRADIENT", "error", where,
                         found="gradient fill inside a chart series",
                         expected="a flat series colour from brand.colorRules.chartSeries",
                         rule="brand.colorRules.chartGradientFillForbidden is true.",
                         fix="Fill each series with one solid colour from %s."
                             % ", ".join((self.brand.get("colorRules") or {})
                                         .get("chartSeries") or []),
                         key=("COLOR.CHART_GRADIENT", ctx.index, rec.path))

        for clr in root.iter("%ssrgbClr" % _A):
            val = clr.get("val")
            if val and bl.is_hex(val):
                self.report_color(where, val, False, "chart")

        for latin in root.iter("%slatin" % _A):
            family = latin.get("typeface")
            if not family:
                continue
            if _norm_family(family) not in self.families:
                self.add("TYPE.OFF_FAMILY", "error", where,
                         found="%s (chart text)" % family,
                         expected=", ".join(sorted(self.families.values())),
                         rule="Chart labels use the brand type families too.",
                         fix="Set the chart text font to %s."
                             % ((self.brand.get("type") or {}).get("family") or "the brand family"),
                         key=("TYPE.OFF_FAMILY", where, _norm_family(family)))

        for def_rpr in root.iter("%sdefRPr" % _A):
            sz = def_rpr.get("sz")
            if sz is None:
                continue
            try:
                size = float(sz) / 100.0
            except (TypeError, ValueError):
                continue
            if size < self.min_body_pt - 1e-6:
                self.add("TYPE.BELOW_MIN", "error", where,
                         found="%.4gpt (chart text)" % size,
                         expected=">= %.4gpt" % self.min_body_pt,
                         rule="brand.type.minBodyPt applies to chart labels.",
                         fix="Raise chart label sizes to at least %.4gpt." % self.min_body_pt,
                         key=("TYPE.BELOW_MIN", where, round(size, 2)))

    # -- structure ----------------------------------------------------------

    def check_structure(self, slide_count):
        # type: (int) -> None
        archetypes = [(c.index, c.archetype, c.archetype_source) for c in self.contexts]

        min_slides = int(self.grammar.rule("minSlides", 3))
        if slide_count < min_slides:
            self.add("STRUCTURE.TOO_FEW_SLIDES", "error", "deck",
                     found="%d slides" % slide_count,
                     expected=">= %d" % min_slides,
                     rule="grammar.deckRules.minSlides.",
                     fix="A deck needs a cover, at least one content slide and a close.")

        known = [a for a in archetypes if a[1]]
        if not known:
            self.add("STRUCTURE.UNKNOWN", "info", "deck",
                     found="no archetype could be established for any slide",
                     expected="an --ir file, or a deck built from the grammar",
                     rule="Structural rules need slide archetypes.",
                     fix="Re-run with --ir pointing at the deck IR that produced this file.")
            return

        from_ir = any(src == "ir" for _i, _a, src in archetypes)
        hard = "error" if from_ir else "warn"
        note = "" if from_ir else " (archetype inferred from geometry; pass --ir to be sure)"

        opener = str(self.grammar.rule("mustOpenWith", "cover"))
        first = archetypes[0][1] if archetypes else None
        if first != opener:
            self.add("STRUCTURE.NO_COVER", hard, "slide 1",
                     found=str(first or "unknown") + note,
                     expected=opener,
                     rule="grammar.deckRules.mustOpenWith.",
                     fix="Make slide 1 a %s archetype." % opener)

        closers = self.grammar.rule("mustCloseWith", ["closing"]) or ["closing"]
        if isinstance(closers, str):
            closers = [closers]
        last = archetypes[-1][1] if archetypes else None
        if last not in closers:
            self.add("STRUCTURE.NO_CLOSING", hard, "slide %d" % slide_count,
                     found=str(last or "unknown") + note,
                     expected=" or ".join(str(c) for c in closers),
                     rule="grammar.deckRules.mustCloseWith.",
                     fix="End the deck with a %s slide carrying one concrete next step."
                         % closers[0])

        max_run = int(self.grammar.rule("maxConsecutiveSameArchetype", 2))
        run_len = 0
        run_id = None
        for idx, aid, _src in archetypes:
            if aid and aid == run_id:
                run_len += 1
            else:
                run_id, run_len = aid, 1
            if aid and run_len == max_run + 1:
                self.add("STRUCTURE.REPEATED_ARCHETYPE", "warn", "slide %d" % idx,
                         found="%d consecutive %s slides start at slide %d"
                               % (run_len, aid, idx - max_run),
                         expected="at most %d in a row" % max_run,
                         rule="grammar.deckRules.maxConsecutiveSameArchetype.",
                         fix="Alternate archetypes: swap one for its mirror "
                             "(text-visual <-> visual-text) or a stats/quote break.",
                         key=("STRUCTURE.REPEATED_ARCHETYPE", aid, idx))

        gap_max = int(self.grammar.rule("sectionBreakEveryNSlidesMax", 8))
        since = 0
        for idx, aid, _src in archetypes:
            if aid == "section-break" or aid in ("cover",):
                since = 0
                continue
            since += 1
            if since > gap_max:
                self.add("STRUCTURE.MISSING_SECTION_BREAK", "warn", "slide %d" % idx,
                         found="%d slides since the last section break" % since,
                         expected="a section break at least every %d slides" % gap_max,
                         rule="grammar.deckRules.sectionBreakEveryNSlidesMax.",
                         fix="Insert a section-break slide before slide %d." % idx,
                         key=("STRUCTURE.MISSING_SECTION_BREAK", idx))
                since = 0

    # -- learned rules ------------------------------------------------------

    def _scoped_text(self, scope):
        # type: (str) -> List[Tuple[str, str]]
        """[(where, text)] for a learned-rule scope."""
        out = []
        for ctx in self.contexts:
            for rec in ctx.recs:
                if not rec.text:
                    continue
                is_eyebrow = role_is_eyebrow(rec, ctx, self.grammar)
                is_title = _is_title_shape(rec)
                if scope == "title" and not is_title:
                    continue
                if scope == "eyebrow" and not is_eyebrow:
                    continue
                if scope == "body" and (is_title or is_eyebrow):
                    continue
                out.append((ctx.where(rec.path), rec.text))
        return out

    def _scoped_runs(self, scope):
        # type: (str) -> List[Tuple[str, Dict[str, Any]]]
        out = []
        for ctx in self.contexts:
            for rec in ctx.recs:
                if not rec.runs:
                    continue
                is_eyebrow = role_is_eyebrow(rec, ctx, self.grammar)
                is_title = _is_title_shape(rec)
                if scope == "title" and not is_title:
                    continue
                if scope == "eyebrow" and not is_eyebrow:
                    continue
                if scope == "body" and (is_title or is_eyebrow):
                    continue
                for para in rec.runs:
                    for run in para["runs"]:
                        if (run["text"] or "").strip():
                            out.append((ctx.where(rec.path), run))
        return out

    def _all_colors(self):
        # type: () -> List[Tuple[str, str]]
        out = []
        for ctx in self.contexts:
            for rec in ctx.recs:
                where = ctx.where(rec.path)
                for info in (rec.fill, rec.line):
                    for hexed in info.get("colors") or []:
                        if hexed:
                            out.append((where, hexed))
                for para in rec.runs:
                    for run in para["runs"]:
                        for hexed in run["fill"].get("colors") or []:
                            if hexed:
                                out.append((where, hexed))
        return out

    def check_learned_rules(self):
        # type: () -> None
        learned = self.brand.get("learnedRules") or {}
        rules = learned.get("rules") if isinstance(learned, dict) else None
        if not isinstance(rules, list):
            return

        for raw in rules:
            if not isinstance(raw, dict):
                continue
            vid = str(raw.get("id") or "LEARNED.RULE")
            severity = str(raw.get("severity") or "warn").lower()
            if severity not in bl.SEVERITIES:
                severity = "warn"
            kind = str(raw.get("kind") or "")
            scope = str(raw.get("scope") or "any")
            rule_text = str(raw.get("rule") or "Learned rule %s." % vid)
            fix_text = str(raw.get("fix") or "Apply the learned correction.")
            value = raw.get("value")
            values = value if isinstance(value, list) else [value]

            try:
                if kind == "forbid_text":
                    needles = [str(v) for v in values if v is not None]
                    for where, text in self._scoped_text(scope):
                        for hit in bl.find_forbidden_phrases(text, needles):
                            self.add(vid, severity, where,
                                     found=hit, expected="copy without %r" % hit,
                                     rule=rule_text, fix=fix_text,
                                     key=(vid, where, hit))

                elif kind == "require_text":
                    needles = [str(v) for v in values if v is not None]
                    haystack = "\n".join(t for _w, t in self._scoped_text(scope)).lower()
                    for needle in needles:
                        if needle.strip() and needle.lower() not in haystack:
                            self.add(vid, severity, "deck (scope=%s)" % scope,
                                     found="%r never appears" % needle,
                                     expected="%r somewhere in scope" % needle,
                                     rule=rule_text, fix=fix_text,
                                     key=(vid, needle))

                elif kind == "regex":
                    for pattern in [str(v) for v in values if v is not None]:
                        try:
                            rx = re.compile(pattern, re.IGNORECASE)
                        except re.error as exc:
                            self.add(vid, "info", "brand rules.local.json",
                                     found="invalid regex %r: %s" % (pattern, exc),
                                     expected="a compilable Python regular expression",
                                     rule=rule_text,
                                     fix="Fix the pattern in rules.local.json.",
                                     key=(vid, "badregex", pattern))
                            continue
                        for where, text in self._scoped_text(scope):
                            m = rx.search(text)
                            if m:
                                self.add(vid, severity, where,
                                         found="matched %r at %r" % (pattern,
                                                                     _snippet(m.group(0))),
                                         expected="copy that does not match %r" % pattern,
                                         rule=rule_text, fix=fix_text,
                                         key=(vid, where, pattern))

                elif kind == "forbid_color":
                    banned = set()
                    for v in values:
                        if bl.is_hex(v):
                            banned.add(bl.normalize_hex(str(v)))
                    for where, hexed in self._all_colors():
                        if bl.normalize_hex(hexed) in banned:
                            self.add(vid, severity, where,
                                     found=bl.normalize_hex(hexed),
                                     expected="any approved brand colour",
                                     rule=rule_text, fix=fix_text,
                                     key=(vid, where, bl.normalize_hex(hexed)))

                elif kind in ("max_font_size", "min_font_size", "forbid_font_size"):
                    numbers = []
                    for v in values:
                        try:
                            numbers.append(float(v))
                        except (TypeError, ValueError):
                            continue
                    if not numbers:
                        continue
                    for where, run in self._scoped_runs(scope):
                        size = run["size"]
                        if size is None:
                            continue
                        size = float(size)
                        hit = False
                        expected = ""
                        if kind == "max_font_size" and size > max(numbers) + 1e-6:
                            hit, expected = True, "<= %.4gpt" % max(numbers)
                        elif kind == "min_font_size" and size < min(numbers) - 1e-6:
                            hit, expected = True, ">= %.4gpt" % min(numbers)
                        elif kind == "forbid_font_size" and any(
                                abs(size - n) < 1e-6 for n in numbers):
                            hit, expected = True, "any size other than %s" % ", ".join(
                                "%.4gpt" % n for n in numbers)
                        if hit:
                            self.add(vid, severity, where,
                                     found="%.4gpt" % size, expected=expected,
                                     rule=rule_text, fix=fix_text,
                                     key=(vid, where, round(size, 2)))

                else:
                    self.add(vid, "info", "brand rules.local.json",
                             found="unsupported rule kind %r" % kind,
                             expected="one of forbid_text, forbid_color, require_text, "
                                      "max_font_size, min_font_size, forbid_font_size, regex",
                             rule=rule_text,
                             fix="Rewrite the rule using a supported kind, or extend "
                                 "validate_deck.py.",
                             key=(vid, "unsupported", kind))

            except Exception as exc:  # a bad learned rule must never abort a run
                self.add(vid, "info", "brand rules.local.json",
                         found="rule raised %s: %s" % (type(exc).__name__, exc),
                         expected="a well-formed learned rule",
                         rule=rule_text,
                         fix="Check the rule's value and scope in rules.local.json.",
                         key=(vid, "raised"))


def _is_title_shape(rec):
    # type: (ShapeRec) -> bool
    """True when a shape holds a slide title rather than title-sized copy."""
    if rec.semantic is not None:
        return rec.semantic == "title"
    return rec.role in TITLE_ROLES


def _norm_family(family):
    # type: (Optional[str]) -> str
    return re.sub(r"[^a-z0-9]+", "", str(family or "").lower())


def _snippet(text, limit=72):
    # type: (str, int) -> str
    flat = re.sub(r"\s+", " ", str(text or "")).strip()
    if len(flat) <= limit:
        return flat
    return flat[:limit - 1] + "…"


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def read_custom_brand(path):
    # type: (str) -> Optional[str]
    """Value of the docProps custom property 'brand-studio.brand', if present."""
    try:
        with zipfile.ZipFile(path) as zf:
            if "docProps/custom.xml" not in zf.namelist():
                return None
            blob = zf.read("docProps/custom.xml")
    except (IOError, OSError, zipfile.BadZipFile, KeyError):
        return None
    try:
        root = ET.fromstring(blob)
    except ET.ParseError:
        return None
    for prop in root.findall("%sproperty" % _CUSTOM_NS):
        if str(prop.get("name") or "") != CUSTOM_PROP_NAME:
            continue
        for child in list(prop):
            if child.text and child.text.strip():
                return child.text.strip()
    return None


class _Parser(argparse.ArgumentParser):
    def error(self, message):
        self.print_usage(sys.stderr)
        sys.stderr.write("%s: error: %s\n" % (self.prog, message))
        raise SystemExit(1)


def build_parser():
    # type: () -> argparse.ArgumentParser
    parser = _Parser(
        prog="validate_deck.py",
        description="Validate a .pptx against a brand-studio brand profile and the "
                    "shared deck grammar.",
        epilog="Exit 0 when there are no errors, 2 when there are (or when --max-warn "
               "is exceeded), 1 on internal failure.")
    parser.add_argument("pptx", metavar="FILE.pptx",
                        help="the presentation to validate")
    parser.add_argument("--brand", metavar="ID", default=None,
                        help="brand id. Default: the docProps custom property "
                             "'%s' if present, else '%s'."
                             % (CUSTOM_PROP_NAME, DEFAULT_BRAND))
    parser.add_argument("--format", dest="fmt", choices=("json", "human"), default="json",
                        help="output format (default: json)")
    parser.add_argument("--ir", metavar="DECK.json", default=None,
                        help="the deck IR the file was built from. Supplying it makes the "
                             "STRUCTURE checks authoritative and sharpens role detection.")
    parser.add_argument("--max-warn", dest="max_warn", metavar="N", type=int, default=None,
                        help="fail the run when the warning count exceeds N")
    return parser


def main(argv=None):
    # type: (Optional[Sequence[str]]) -> int
    parser = build_parser()
    args = parser.parse_args(argv)

    target = os.path.abspath(args.pptx)
    if not os.path.isfile(target):
        sys.stderr.write("validate_deck.py: no such file: %s\n" % target)
        return 1
    if args.max_warn is not None and args.max_warn < 0:
        sys.stderr.write("validate_deck.py: --max-warn must be >= 0\n")
        return 1

    brand_id = args.brand or read_custom_brand(target) or DEFAULT_BRAND
    resolved = bl.resolve_brand(brand_id)
    if not resolved.get("brand"):
        sys.stderr.write(
            "validate_deck.py: unknown brand %r. Candidates: %s\n"
            % (brand_id, json.dumps(resolved.get("candidates") or [])))
        return 1

    try:
        brand = bl.load_brand(resolved["brand"])
        grammar = Grammar(bl.load_grammar())
    except (bl.BrandNotFound, ValueError, IOError) as exc:
        sys.stderr.write("validate_deck.py: %s\n" % exc)
        return 1

    ir = None
    if args.ir:
        ir_path = os.path.abspath(args.ir)
        if not os.path.isfile(ir_path):
            sys.stderr.write("validate_deck.py: no such IR file: %s\n" % ir_path)
            return 1
        try:
            with open(ir_path, "r") as fh:
                ir = json.load(fh)
        except ValueError as exc:
            sys.stderr.write("validate_deck.py: %s is not valid JSON: %s\n" % (ir_path, exc))
            return 1
        if not isinstance(ir, dict):
            sys.stderr.write("validate_deck.py: %s must contain a JSON object\n" % ir_path)
            return 1

    try:
        validator = DeckValidator(target, brand, grammar, ir)
        report = validator.run()
    except Exception as exc:  # noqa: BLE001 - any failure here is an internal failure
        sys.stderr.write("validate_deck.py: internal failure on %s: %s: %s\n"
                         % (target, type(exc).__name__, exc))
        return 1

    if args.max_warn is not None:
        warns = report.counts()["warn"]
        if warns > args.max_warn:
            report.add("STRUCTURE.WARN_BUDGET", severity="error", where="deck",
                       found="%d warnings" % warns,
                       expected="at most %d (--max-warn)" % args.max_warn,
                       rule="The caller set a warning budget for this run.",
                       fix="Clear warnings until the count is %d or lower." % args.max_warn)

    payload = report.to_json(target=target, brand=brand.get("id", ""), kind=KIND)
    if args.fmt == "human":
        sys.stdout.write(report.to_human(target=target, brand=brand.get("id", ""), kind=KIND))
    else:
        sys.stdout.write(json.dumps(payload, indent=2) + "\n")
    return 2 if payload["counts"]["error"] > 0 else 0


if __name__ == "__main__":
    sys.exit(main())
